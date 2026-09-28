from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = "v0.2.0-shopping-list.docx"

items = [
    ("Wired numeric keypad", 1000, 1, "LiteLuck Enterprises", "https://jiji.co.ke/nairobi-central/computer-accessories/wired-numeric-keypad-hnfY7tyK6hT5BIw1vHKeIQK0.html?page=1&pos=1&cur_pos=1&ads_per_page=10&lid=mPfxkidRCy-5CIQv"),
    ("Type C to USB 3.0 adapter", 500, 2, "Jeffry Tech", "https://jiji.co.ke/nairobi-central/computer-accessories/type-c-to-usb3-0-data-transfer-adapter-male-to-female-plug-c-52n2X5VP0QxVuQIkfZTGXvKv.html?page=1&pos=2&cur_pos=2&ads_per_page=10&lid=Hly6khU0dvxQD1c2"),
    ("Assorted crimp wire terminal kit", 3000, 1, "Weiss imports", "https://jiji.co.ke/ngara/car-parts-and-accessories/assorted-crimp-wire-terminal-kit-vXkBUZvKRz9nMbYhxeKzeKZG.html?page=1&pos=3&cur_pos=3&ads_per_page=10&lid=_0fSlSvBnWGbm97t"),
    ("Lead free solder paste flux", 500, 1, "santaecommerce.com", "https://jiji.co.ke/nairobi-central/hand-and-power-tools/158degc-lead-free-solder-paste-high-performance-smd-flux-low-temp-1BSBl9FaeeMiFhoq8FopDlnL.html?page=1&pos=4&cur_pos=4&ads_per_page=10&lid=SlEaX4pYbCfD2kxf"),
    ("Gold 3.5 mm audio jack connector", 300, 1, "Jeffry Tech", "https://jiji.co.ke/nairobi-central/cell-phones-tablets-accessories/3-5mm-3-pole-stereo-audio-male-jack-plug-connector-solder-pa-rsYSEq0XXtJ1kOMgFtNvz6zJ.html?page=1&pos=5&cur_pos=5&ads_per_page=10&lid=l0Lzfmsc2Yja8bpZ"),
    ("Asahi 1.0 mm 100 g solder wire", 999, 1, "Budget . World", "https://jiji.co.ke/nairobi-central/power-tools/smooth-flow-strong-joints-asahi-solder-wire-1-0mm-100g-rosin-core-mLWAGTHpEM5UMnJbidpBlMlr.html?page=1&pos=8&cur_pos=8&ads_per_page=10&lid=ytjO-LbIeX43HO0W"),
    ("14 piece 60 W soldering kit", 1997, 1, "Edison Supplies Kenya", "https://jiji.co.ke/nairobi-central/power-tools/14pcs-soldering-iron-60w-kit-tips-solder-gun-sucker-tweezers-wire-set-doQ7Ni06xPt4mV4RX7wqpV3h.html?page=1&pos=9&cur_pos=9&ads_per_page=10&lid=jeP-VU3bcRi8eVdm"),
    ("Silver 3.5 mm audio jack connector", 500, 1, "Jeffry Tech", "https://jiji.co.ke/nairobi-central/computer-accessories/3-5mm-male-mini-jack-stereo-solder-plug-vWgiYYN7UndFxWXmbTDjhuGb.html?page=1&pos=10&cur_pos=10&ads_per_page=10&lid=1u5WaljLUjWBZW31"),
    ("Ugreen 3.5 mm aux extension cable", 1600, 2, "Jeffry Tech", "https://jiji.co.ke/nairobi-central/computer-accessories/ugreen-3-5mm-male-to-female-extension-cable-with-mic-support-dbJUBwFfuDynFYfZLd3LGMiP.html?page=2&pos=3&cur_pos=3&ads_per_page=10&lid=t3uA_F8KiyjY_0_V"),
    ("3 ft USB 3.0 extension cable", 999, 4, "Jeffry Tech", "https://jiji.co.ke/nairobi-central/computer-accessories/3ft-usb-3-0-high-speed-superspeed-extension-cable-a-male-to-a-female-d5erAuuQWmXBfJCZby9DDktd.html?page=2&pos=5&cur_pos=5&ads_per_page=10&lid=XG0OIMUjpg25mir5"),
    ("USB C male to female extension cable", 999, 1, "Jeffry Tech", "https://jiji.co.ke/nairobi-central/computer-accessories/usb-3-1-type-c-male-to-female-extension-cable-usb-c-m-f-extender-cord-s1jTbXAUuNHRDEEVV8X83qp.html?page=2&pos=6&cur_pos=6&ads_per_page=10&lid=dws-rq22pqJZaapn"),
    ("Ugreen CM420 powered USB hub", 4999, 1, "Daytech Accessories", "https://jiji.co.ke/nairobi-central/computer-accessories/ugreen-cm420-usb-hub-7-ports-usb-3-0-utHVXFgQvIhONHYwZgF7xxmv.html?page=2&pos=7&cur_pos=7&ads_per_page=10&lid=dRhEdsjzNa2RbGpY"),
    ("Ugreen CM264 USB card reader", 999, 1, "Vertex Stores", "https://jiji.co.ke/nairobi-central/computer-accessories/ugreen-cm264-usb-3-0-multifunctional-card-reader-tf-sd-card-reader-18U066jJHdkNVB3zzZmIVOf8.html?page=1&pos=16&cur_pos=16&ads_per_page=24&ads_count=24&lid=QDhnl_X74r38UE0i&indexPosition=15"),
    ("Ugreen RCA to 3.5 mm stereo cable 1 m", 1928, 2, "Jeffry Tech", "https://jiji.co.ke/nairobi-central/computer-accessories/ugreen-rca-to-3-5mm-cable-phono-mini-jack-stereo-1m-APQ180J0P9PKxxgT93FyVHEU.html?page=4&pos=1&cur_pos=1&ads_per_page=8&lid=j74bTW1CdNc-eYlR"),
    ("RCA male to double female splitter", 650, 2, "Absolute car accessories", "https://jiji.co.ke/ngara/accessories-and-supplies-for-electronics/rca-split-cable-nJNUwlISlQK8KtjtZIMMz7cp.html?page=1&pos=2&cur_pos=2&ads_per_page=15&ads_count=134&lid=jHCd-XZevieyaj23&indexPosition=30"),
    ("10 m twin RCA phono audio cable male", 899, 1, "Jeffry Tech", "https://jiji.co.ke/nairobi-central/accessories-and-supplies-for-electronics/10m-twin-red-white-2-x-rca-phono-audio-left-right-cable-male-CbVHQIA488a0cvY6l4RptpD.html?page=2&pos=12&cur_pos=12&ads_per_page=24&ads_count=298&lid=nX_VUGgLtidN0VbP&indexPosition=35"),
]

contacts = {
    "LiteLuck Enterprises": "0712 653 679",
    "Jeffry Tech": "0727 470 163 / 0720 744 050",
    "Weiss imports": "+254 708 365 111",
    "santaecommerce.com": "0747 047 470",
    "Budget . World": "+254 115 429 467",
    "Edison Supplies Kenya": "0743 895 315 / 0115 875 133",
    "Daytech Accessories": "+254 710 694 608",
    "Vertex Stores": "+254 718 035 859",
    "Absolute car accessories": "+254 796 757 863",
}

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)

def borders(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), "D9D9D9")

def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + name))
        if node is None:
            node = OxmlElement("w:" + name)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")

def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)

def hyperlink(paragraph, label, url):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    run.append(r_pr)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.65)
section.right_margin = Inches(0.65)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].font.color.rgb = RGBColor(0, 0, 0)
for name in ("Title", "Heading 1", "Heading 2"):
    styles[name].font.name = "Aptos Display"
    styles[name].font.color.rgb = RGBColor(0, 0, 0)
styles["Title"].font.size = Pt(22)
styles["Heading 1"].font.size = Pt(14)
styles["Heading 2"].font.size = Pt(10.5)
styles["Heading 2"].paragraph_format.space_before = Pt(4)
styles["Heading 2"].paragraph_format.space_after = Pt(0)
styles["List Bullet"].font.name = "Aptos"
styles["List Bullet"].font.size = Pt(9.25)
styles["List Bullet"].paragraph_format.space_before = Pt(0)
styles["List Bullet"].paragraph_format.space_after = Pt(0)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.add_run("V0 2 0 Shopping List")

intro = doc.add_paragraph()
intro.paragraph_format.space_after = Pt(8)
intro.add_run("Prepared 4 September 2026. ").bold = True
intro.add_run("This list records the prices shown on the supplied Jiji listings and the requested quantities. The purchase total is KSh 29,544 before delivery or any vendor-negotiated changes.")

doc.add_paragraph("Items and pricing", style="Heading 1")
table = doc.add_table(rows=1, cols=5)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [Inches(3.0), Inches(1.1), Inches(0.65), Inches(1.15), Inches(1.2)]
headers = ["Item", "Unit price", "Quantity", "Total price", "Vendor"]
for cell, width, text in zip(table.rows[0].cells, widths, headers):
    cell.width = width
    shade(cell, "17365D")
    borders(cell)
    set_cell_margins(cell)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if text != "Item" else WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    r.bold = True
    r.font.color.rgb = RGBColor(255, 255, 255)
    r.font.size = Pt(9)
set_repeat_table_header(table.rows[0])

total = 0
for idx, (item, unit, qty, vendor, url) in enumerate(items):
    row = table.add_row()
    line_total = unit * qty
    total += line_total
    values = [item, f"KSh {unit:,.0f}", str(qty), f"KSh {line_total:,.0f}", vendor]
    for col, (cell, width, value) in enumerate(zip(row.cells, widths, values)):
        cell.width = width
        borders(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        if idx % 2 == 1:
            shade(cell, "F3F7FB")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col in (1, 2, 3) else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(value)
        r.font.size = Pt(9)

grand = table.add_row()
grand.cells[0].merge(grand.cells[3])
for c in grand.cells:
    shade(c, "DCE6F1")
    borders(c)
    set_cell_margins(c)
grand.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
gr = grand.cells[0].paragraphs[0].add_run("Complete total price")
gr.bold = True
grand.cells[4].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
gr = grand.cells[4].paragraphs[0].add_run(f"KSh {total:,.0f}")
gr.bold = True

doc.add_paragraph("Vendors and contacts", style="Heading 1")
note = doc.add_paragraph()
note.paragraph_format.space_after = Pt(6)
for run in note.runs:
    run.font.size = Pt(9.5)
note.add_run("Contact note. ").bold = True
note.add_run("The contacts below combine the supplied vendor phone numbers with the live Jiji item links.")
for run in note.runs:
    run.font.size = Pt(9.5)

vendor_order = []
for item in items:
    if item[3] not in vendor_order:
        vendor_order.append(item[3])

for vendor in vendor_order:
    p = doc.add_paragraph(style="Heading 2")
    vendor_run = p.add_run(vendor)
    vendor_run.bold = True
    p.add_run("  Contact: ")
    label = p.runs[-1]
    label.bold = True
    contact = p.add_run(contacts[vendor])
    p.paragraph_format.space_after = Pt(1)
    vendor_run.font.size = Pt(9.75)
    label.bold = True
    label.font.size = Pt(9.5)
    contact.font.size = Pt(9.5)
    for item, unit, qty, item_vendor, url in items:
        if item_vendor == vendor:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(0)
            hyperlink(p, item, url)
            p.add_run(f" - {qty}")

doc.save(OUT)
print(f"Created {OUT} with {len(items)} items and total KSh {total:,.0f}")
